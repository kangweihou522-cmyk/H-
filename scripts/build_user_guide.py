"""Build matching offline HTML and Word guides from docs/USER_GUIDE.md.

Run with --check in CI to reject stale artifacts without changing any files.
Documentation tools are separate from the workstation's runtime dependencies.
"""
from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
from html import escape
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile, ZipInfo, ZIP_STORED

try:
    import markdown
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
except ImportError as error:
    raise SystemExit('缺少说明书构建工具，请运行 python -m pip install -r docs/requirements.txt') from error

ROOT = Path(__file__).resolve().parent.parent
SOURCE = Path('docs/USER_GUIDE.md')
OUTPUTS = (Path('static/guide/index.html'), Path('docs/USER_GUIDE.docx'), Path('docs/guide-manifest.json'))
RUNTIME_PATTERNS = ('*.py', '*.sh', '*.bat', 'Dockerfile*', 'requirements.txt', 'README.md',
                    'scripts/*.py', 'scripts/*.R', 'scripts/*.sh', 'windows/*.py',
                    'windows/*.sh', 'windows/*.ps1', 'static/*.html', 'static/*.js',
                    'static/*.css', 'static/templates/*', 'docs/WORKFLOW.md', 'docs/WINDOWS_D_DRIVE.md')


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def normalized(path: Path) -> bytes:
    # Git archives use LF for PS1 while a Windows checkout may use CRLF.
    return path.read_bytes().replace(b'\r\n', b'\n')


def source_fingerprint(root: Path) -> dict[str, str]:
    paths = {p for pattern in RUNTIME_PATTERNS for p in root.glob(pattern) if p.is_file()}
    paths.discard(root/'scripts/build_user_guide.py')
    return {p.relative_to(root).as_posix(): sha(normalized(p)) for p in sorted(paths)}


def heading_id(value, separator):
    return 'section-'+sha(value.encode('utf-8'))[:12]


class Node:
    def __init__(self, tag='', attrs=()):
        self.tag, self.attrs, self.children = tag, dict(attrs), []

    def text(self):
        return ''.join(n.text() if isinstance(n, Node) else n for n in self.children)

    def descendants(self, tag):
        for child in self.children:
            if isinstance(child, Node):
                if child.tag == tag:
                    yield child
                yield from child.descendants(tag)


class Tree(HTMLParser):
    def __init__(self, body):
        super().__init__(convert_charrefs=True)
        self.root = Node('root')
        self.stack = [self.root]
        self.feed(body)

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs)
        self.stack[-1].children.append(node)
        if tag not in {'br', 'hr', 'img', 'input', 'meta', 'link'}:
            self.stack.append(node)

    def handle_endtag(self, tag):
        for index in range(len(self.stack)-1, 0, -1):
            if self.stack[index].tag == tag:
                self.stack = self.stack[:index]
                break

    def handle_data(self, value):
        self.stack[-1].children.append(value)


def element(tag, **attrs):
    node = OxmlElement('w:'+tag)
    for key, value in attrs.items():
        node.set(qn('w:'+key), str(value))
    return node


def link(paragraph, text, target, internal=False):
    item = element('hyperlink', anchor=target) if internal else element('hyperlink')
    if not internal:
        item.set(qn('r:id'), paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True))
    run = element('r')
    properties = element('rPr')
    properties.append(element('color', val='126657'))
    properties.append(element('u', val='single'))
    run.append(properties)
    value = element('t')
    value.text = text
    run.append(value)
    item.append(run)
    paragraph._p.append(item)


def inline(paragraph, children, bold=False, italic=False, code=False):
    for child in children:
        if isinstance(child, str):
            run = paragraph.add_run(child)
            run.bold, run.italic = bold or None, italic or None
            if code:
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run._r.get_or_add_rPr().append(element('shd', fill='EDF4EF'))
        elif child.tag == 'a':
            target = child.attrs.get('href', '')
            link(paragraph, child.text(), target.lstrip('#') if target.startswith('#') else target, internal=target.startswith('#'))
        elif child.tag == 'br':
            paragraph.add_run().add_break()
        else:
            inline(paragraph, child.children, bold or child.tag in {'strong', 'b'},
                   italic or child.tag in {'em', 'i'}, code or child.tag == 'code')


def make_word(tree, meta, images=None):
    images = images or {}
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.1), Cm(2)
    section.left_margin, section.right_margin = Cm(2.1), Cm(2.1)
    section.header_distance = section.footer_distance = Cm(1)
    for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Paragraph']:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '微软雅黑')
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 27), ('Heading 1', 19), ('Heading 2', 14), ('Heading 3', 11.5)]:
        style = doc.styles[name]
        style.font.size, style.font.color.rgb = Pt(size), RGBColor.from_string('126657')
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    header.text = '小鼠 bulk RNA-seq 工作站  /  使用说明与方法参考'
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('说明书 '+meta['guide_version']+'  ·  ')
    footer._p.append(element('fldSimple', instr='PAGE'))
    properties = doc.core_properties
    properties.title, properties.subject = meta['title'], '操作流程、统计原理、结果解读与参考文献'
    properties.author = properties.last_modified_by = 'Mouse RNA-seq Workstation'
    properties.created = properties.modified = datetime.strptime(meta['updated'], '%Y-%m-%d')
    properties.keywords = 'Mus musculus, RNA-seq, Salmon, DESeq2, GO, KEGG, GSEA, WGCNA'
    doc.settings.element.append(element('updateFields', val='true'))
    doc.add_paragraph('LAB WORKSTATION  /  METHODS & USER GUIDE', style='Subtitle')
    doc.add_paragraph(meta['title'], style='Title')
    doc.add_paragraph('从RNA与建库测序讲起，循序理解表达、统计与工作站操作', style='Subtitle')
    doc.add_paragraph(f"说明书版本：{meta['guide_version']}\n适用程序：v{meta['application_version']}\n更新日期：{meta['updated']}")
    doc.add_paragraph('网页使用说明与本 Word 文档由同一份文稿生成。实验结论应结合设计、质量控制、效应量及独立验证。')
    doc.add_paragraph('程序内容指纹：'+meta['code_fingerprint'], style='Caption')
    doc.add_page_break()
    doc.add_paragraph('阅读目录', style='Heading 1')
    doc.add_paragraph('目录只列主要单元；章内小节可在Word“视图 → 导航窗格”中查找。')
    headings = [n for n in tree.root.children if isinstance(n, Node) and n.tag == 'h2']
    for heading in headings:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(.5 if heading.tag == 'h3' else 0)
        paragraph.paragraph_format.space_after = Pt(2)
        link(paragraph, heading.text(), heading.attrs['id'], internal=True)
    doc.add_page_break()
    bookmark = 1

    def blocks(nodes, depth=0):
        nonlocal bookmark
        for node in nodes:
            if not isinstance(node, Node):
                continue
            tag = node.tag
            if tag == 'h1':
                continue  # The title is on the cover.
            if tag in {'h2', 'h3', 'h4', 'h5', 'h6'}:
                paragraph = doc.add_paragraph(style='Heading '+str(min(3, int(tag[1])-1)))
                paragraph._p.append(element('bookmarkStart', id=bookmark, name=node.attrs.get('id', 'heading-'+str(bookmark))))
                inline(paragraph, node.children)
                paragraph._p.append(element('bookmarkEnd', id=bookmark))
                bookmark += 1
            elif tag == 'p':
                pictures = list(node.descendants('img'))
                for picture in pictures:
                    data = images[picture.attrs['src']]
                    shape = doc.add_picture(io.BytesIO(data), width=Cm(16))
                    shape._inline.docPr.set('descr', picture.attrs.get('alt', ''))
                    if picture.attrs.get('alt'):
                        doc.add_paragraph(picture.attrs['alt'], style='Caption')
                if node.text().strip():
                    inline(doc.add_paragraph(), node.children)
            elif tag == 'pre':
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.line_spacing = 1
                paragraph._p.get_or_add_pPr().append(element('shd', fill='EDF4EF'))
                run = paragraph.add_run(node.text().rstrip('\n'))
                run.font.name, run.font.size = 'Consolas', Pt(8)
            elif tag in {'ul', 'ol'}:
                for index, item in enumerate(n for n in node.children if isinstance(n, Node) and n.tag == 'li'):
                    paragraph = doc.add_paragraph(style='List Paragraph')
                    paragraph.paragraph_format.left_indent = Cm(.5*(depth+1))
                    paragraph.paragraph_format.first_line_indent = Cm(-.35)
                    paragraph.add_run((str(index+1)+'. ') if tag == 'ol' else '• ')
                    inline(paragraph, [n for n in item.children if not isinstance(n, Node) or n.tag not in {'ul', 'ol'}])
                    blocks([n for n in item.children if isinstance(n, Node) and n.tag in {'ul', 'ol'}], depth+1)
            elif tag == 'table':
                rows = list(node.descendants('tr'))
                if not rows:
                    continue
                cells = [[n for n in row.children if isinstance(n, Node) and n.tag in {'td', 'th'}] for row in rows]
                table = doc.add_table(rows=0, cols=max(map(len, cells)))
                table.style = 'Table Grid'
                for index, row in enumerate(cells):
                    target = table.add_row()
                    if index == 0:
                        target._tr.get_or_add_trPr().append(element('tblHeader'))
                    for cell, source in zip(target.cells, row):
                        cell._tc.get_or_add_tcPr().append(element('shd', fill='DCECE5' if index == 0 else ('F4F7F5' if index % 2 == 0 else 'FFFFFF')))
                        inline(cell.paragraphs[0], source.children, bold=index == 0)
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(9)
                    target._tr.get_or_add_trPr().append(element('cantSplit'))
                doc.add_paragraph()
            elif tag == 'blockquote':
                paragraph = doc.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(.5)
                inline(paragraph, node.children, italic=True)
            elif tag != 'hr':
                blocks(node.children, depth)

    blocks(tree.root.children)
    stream = io.BytesIO()
    doc.save(stream)
    # Deterministic package metadata and uncompressed ZIP avoid zlib/version drift.
    result = io.BytesIO()
    with ZipFile(stream) as source, ZipFile(result, 'w', compression=ZIP_STORED) as target:
        for name in sorted(source.namelist()):
            info = ZipInfo(name, date_time=(2020, 1, 1, 0, 0, 0))
            info.compress_type, info.external_attr = ZIP_STORED, 0o600 << 16
            target.writestr(info, source.read(name))
    return result.getvalue()


CSS = '''
:root{--ink:#183531;--green:#126657;--paper:#f4f7f5;--line:#d4e0d9;--muted:#557269}
*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:20px}body{margin:0;color:var(--ink);background:var(--paper);font:15px/1.9 system-ui,"Microsoft Yahei",sans-serif}a{color:var(--green);overflow-wrap:anywhere}header{background:#183c34;color:white;padding:38px max(24px,calc((100vw - 1250px)/2));}header small{color:#c2dfd0;letter-spacing:.1em}h1{font-size:30px;line-height:1.5;margin:10px 0}header p{color:#d4e8df}header a{color:white;display:inline-block;border:1px solid #a9cabd;padding:6px 15px;border-radius:6px;text-decoration:none}.layout{display:grid;grid-template-columns:265px minmax(0,1fr);gap:30px;max-width:1300px;padding:28px 24px;margin:auto}aside{position:sticky;top:20px;align-self:start;max-height:94vh;overflow:auto;font-size:13px}aside ul{list-style:none;padding-left:14px}aside>nav>.toc>ul{padding-left:0}aside a{text-decoration:none;display:block;margin:5px 0}input{width:100%;border:1px solid #b9cbc0;border-radius:6px;font:inherit;padding:10px;background:white}#search-status{font-size:12px;color:var(--muted)}#search-results a{font-size:13px;margin:8px 0}article figure{margin:20px 0}article img{display:block;max-width:100%;height:auto;margin:auto;border:1px solid var(--line);border-radius:8px}figcaption{font-size:12px;color:var(--muted);margin:8px 0}section{background:white;border:1px solid var(--line);border-radius:12px;padding:24px 30px;margin-bottom:22px;min-width:0}h2{font-size:23px;margin-top:4px;line-height:1.6}h3{font-size:18px;margin-top:28px}h4{font-size:16px}p,li{overflow-wrap:anywhere}table{width:100%;border-collapse:collapse;font-size:13px}th,td{text-align:left;vertical-align:top;border:1px solid var(--line);padding:9px;min-width:100px}th{background:#edf4ef}tr:nth-child(even){background:#fafcfb}.table-scroll{overflow-x:auto;margin:18px 0}pre{background:#102a23;color:#d1e9dd;padding:16px;border-radius:8px;overflow-x:auto;font:12px/1.7 Consolas,monospace}code{background:#edf4ef;border-radius:3px;padding:1px 4px;font: .9em Consolas,monospace}pre code{background:transparent;padding:0;color:inherit}blockquote{margin:18px 0;border-left:4px solid #64a28b;background:#edf5ef;padding:8px 18px;color:#365c4e}.stamp{font-size:12px;color:var(--muted);padding:0 4px;overflow-wrap:anywhere}button{font:inherit;background:var(--green);color:white;border:0;border-radius:6px;padding:8px 12px;cursor:pointer}#search-results:empty{display:none}[hidden]{display:none!important}@media(max-width:850px){header{padding:24px}h1{font-size:25px}.layout{display:block;padding:16px}aside{position:static;max-height:none;margin-bottom:20px}aside nav{max-height:220px;overflow:auto;border-bottom:1px solid var(--line)}section{padding:20px 16px}h2{font-size:20px}}@media print{body{background:white;font-size:10pt}.layout{display:block;padding:0}aside,header a,button{display:none}header{background:white;color:#183531;padding:0}header p,header small{color:#183531}section{border:0;padding:0;break-inside:auto}h2,h3,h4{break-after:avoid}table{font-size:8pt}thead{display:table-header-group}tr{break-inside:avoid}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f7f5;color:#183531}.table-scroll{overflow:visible}}
'''
JS = '''
const input=document.getElementById('guide-search'),result=document.getElementById('search-results'),status=document.getElementById('search-status');
const chapters=[...document.querySelectorAll('article section')].map(s=>({heading:s.querySelector('h2'),text:s.textContent.toLocaleLowerCase()})).filter(s=>s.heading);
input.addEventListener('input',()=>{const q=input.value.trim().toLocaleLowerCase();result.replaceChildren();if(!q){status.textContent='输入关键词可定位到相关章节';return;}const found=chapters.filter(s=>s.text.includes(q));status.textContent=`找到 ${found.length} 个相关章节`;for(const s of found){const a=document.createElement('a');a.textContent=s.heading.textContent;a.href='#'+s.heading.id;result.append(a);}});
for(const a of document.querySelectorAll('article a[href^="http"]')){a.target='_blank';a.rel='noopener noreferrer';}
if(location.protocol==='file:')document.getElementById('word-download').href='../../docs/USER_GUIDE.docx';
'''


def make_html(body, toc, meta):
    body = re.sub(r'<h1\b[^>]*>.*?</h1>\s*', '', body, count=1, flags=re.S)
    def figure(match):
        image = next(Tree(match.group(1)).root.descendants('img'))
        return '<figure>'+match.group(1)+'<figcaption>'+escape(image.attrs.get('alt', ''))+'</figcaption></figure>'
    body = re.sub(r'<p>(<img\b[^>]*>)</p>', figure, body)
    sections = ''.join('<section>' +part+'</section>' for part in re.split(r'(?=<h2\b)', body) if part.strip())
    sections = sections.replace('<table>', '<div class="table-scroll"><table>').replace('</table>', '</table></div>')
    title, version, updated = (escape(meta[key]) for key in ['title', 'guide_version', 'updated'])
    return f'''<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{title}</title><style>{CSS}</style></head>
<body><header><small>LAB WORKSTATION · METHODS &amp; USER GUIDE</small><h1>{title}</h1><p>基础知识 · 图解与算例 · 分析原理 · 实操与参考</p><p>说明书 {version} · 适用程序 v{meta['application_version']} · 更新 {updated}</p><a id="word-download" href="/api/manual/word">下载同版本 Word 说明书</a></header>
<div class="layout"><aside><label for="guide-search">搜索说明书</label><input id="guide-search" type="search" placeholder="例如：火山图、KEGG、样本表"><p id="search-status" role="status">输入关键词可定位到相关章节</p><div id="search-results"></div><nav aria-label="说明书目录">{toc}</nav><p class="stamp">程序内容指纹<br>{escape(meta['code_fingerprint'])}</p></aside><article>{sections}<p class="stamp">网页、Word 与文稿同步生成；内容指纹 {meta['manual_sha256']}</p></article></div><script>{JS}</script></body></html>
'''.encode('utf-8')


def build(root: Path = ROOT) -> dict[Path, bytes]:
    source = normalized(root/SOURCE)
    text = source.decode('utf-8-sig')
    edition = re.search(r'说明书版本[：:]\s*(\d{4}\.\d{2}\.\d{2}(?:-r\d+)?)', text)
    if not edition:
        raise ValueError('文稿必须包含说明书版本：YYYY.MM.DD')
    version = edition.group(1)
    files = source_fingerprint(root)
    processor = markdown.Markdown(extensions=['tables', 'fenced_code', 'toc', 'sane_lists'],
                                  extension_configs={'toc': {'slugify': heading_id, 'toc_depth': '2-3'}})
    body = processor.convert(text)
    tree = Tree(body)
    images = {}
    for picture in tree.root.descendants('img'):
        src = picture.attrs.get('src', '')
        asset = (root/'docs'/src).resolve()
        if not src.startswith('assets/') or not asset.is_relative_to((root/'docs/assets').resolve()):
            raise ValueError('教学图必须是docs/assets内的本地文件：'+src)
        if asset.suffix.lower() not in {'.png', '.jpg', '.jpeg'} or not asset.is_file():
            raise ValueError('教学图缺失或不是Word支持的图片：'+src)
        images[src] = asset.read_bytes()
    headings = list(tree.root.descendants('h2'))
    meta = dict(schema=1, title=text.splitlines()[0].lstrip('# '), guide_version=version,
                updated=version.split('-')[0].replace('.', '-'), application_version='2.0',
                manual_sha256=sha(source), source_files=files, assets={name:sha(data) for name,data in images.items()},
                code_fingerprint='sha256:'+sha(json.dumps(files, sort_keys=True).encode()),
                builder_sha256=sha(normalized(Path(__file__))), chapter_count=len(headings),
                reference_links=len({n.attrs['href'] for n in tree.root.descendants('a') if n.attrs.get('href', '').startswith('https://')}),
                artifacts={'source': SOURCE.as_posix(), 'html': OUTPUTS[0].as_posix(), 'word': OUTPUTS[1].as_posix()})
    html = make_html(body, processor.toc, meta)
    word = make_word(tree, meta, images)
    meta['html_sha256'], meta['word_sha256'] = sha(html), sha(word)
    return {OUTPUTS[0]: html, OUTPUTS[1]: word,
            OUTPUTS[2]: (json.dumps(meta, ensure_ascii=False, sort_keys=True, indent=2)+'\n').encode('utf-8'),
            **{Path('static/guide')/name:data for name,data in images.items()}}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Check committed artifacts without changing files')
    args = parser.parse_args()
    artifacts = build()
    changed = [p for p, content in artifacts.items() if not (ROOT/p).is_file() or (ROOT/p).read_bytes() != content]
    if args.check:
        if changed:
            print('使用说明未同步：'+', '.join(map(str, changed)), file=sys.stderr)
            print('核对 docs/USER_GUIDE.md 后运行 python scripts/build_user_guide.py，并一并提交网页、Word和版本清单。', file=sys.stderr)
            return 1
        print('PASS: 使用说明、网页、Word与当前程序指纹一致。')
        return 0
    for path, data in artifacts.items():
        (ROOT/path).parent.mkdir(parents=True, exist_ok=True)
        (ROOT/path).write_bytes(data)
    print('已同步生成网页、Word和版本清单。')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
